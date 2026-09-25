"""Crawl recent menu posts: sitemap -> new menu posts -> post page -> image/PDF files.

Menus come as body images, or attached PDF/Word/Excel files. PDFs are rendered to
page images; .docx/.xlsx are converted to text. Legacy .doc/.xls are skipped for now.

State is a local JSONL file for now (data/raw_posts.jsonl); it moves to Supabase later.

Usage: python -m pipeline.crawl [--schools code1 code2 ...] [--status regular active]
                                [--since-days 14] [--no-download]
"""
import argparse
import asyncio
import csv
import io
import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import unquote

from bs4 import BeautifulSoup

from pipeline import config
from pipeline.fetch import Fetcher, site_url
from pipeline.survey import POST_RE, is_menu_post, parse_sitemap, post_id

STATE_FILE = config.DATA_DIR / "raw_posts.jsonl"

PUBLISHED_RE = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4}),?\s*(\d{1,2}):(\d{2})")
TRAY_SLUG_RE = re.compile(r"khay-an|hinh-anh|bua-an")
IMG_TAG_RE = re.compile(r"<img\b[^>]*>", re.I)
IMG_SRC_RE = re.compile(r'src="(https?://file\w*\.hcm\.shieldix\.app/(uploadimages/(?:news|haydung)|data/doc)/[^"]+)"', re.I)
# Attachments, also found url-encoded inside Office viewer iframes.
DOC_RE = re.compile(r"https?://file\w*\.hcm\.shieldix\.app/data/doc/[^\"'&<>\s]+?\.(?:pdf|docx?|xlsx?)", re.I)
# The related-posts list after the article uses these thumbnail classes.
RELATED_MARKERS = ('class="item-ratio"', 'class="item_hint', 'class="line-5 big_item_news')
MAX_PDF_PAGES = 4


def full_size(url: str) -> str:
    """CDN resizes with ?w=...; the original is much better for OCR."""
    return re.sub(r"\?w=\d+$", "", url.strip())


def article_html(html: str, title: str) -> str:
    """Slice from the article title to the related-posts list, so sidebars don't leak in."""
    body_start = max(html.find("<body"), 0)
    start = html.find(title, body_start) if title else -1
    start = html.rfind("<h1", body_start, start) if start >= 0 else html.find("<h1", body_start)
    start = max(start, body_start)
    ends = [i for i in (html.find(m, start) for m in RELATED_MARKERS) if i >= 0]
    return html[start: min(ends) if ends else len(html)]


def parse_post(html: str, url: str) -> dict:
    soup = BeautifulSoup(html, "lxml")
    og = soup.find("meta", property="og:title")
    h1s = soup.find_all("h1")
    title = " ".join((og.get("content") if og else (h1s[-1].get_text() if h1s else "")).split())

    body = article_html(html, title)
    published = None
    m = PUBLISHED_RE.search(body)
    if m:
        d, mo, y, hh, mm = map(int, m.groups())
        published = datetime(y, mo, d, hh, mm)

    images = []
    for tag in IMG_TAG_RE.findall(body):
        m = IMG_SRC_RE.search(tag)
        if not m:
            continue
        # The shared image library ("haydung") also holds logos and generic
        # featured images; only trust it for images placed in the article text.
        if m.group(2).lower().endswith("haydung") and "anhnoidung" not in tag:
            continue
        if full_size(m.group(1)) not in images:
            images.append(full_size(m.group(1)))

    docs = []
    for u in DOC_RE.findall(unquote(body)):
        if u not in docs:
            docs.append(u)

    slug = POST_RE.search(url).group(2)
    kind = "tray" if TRAY_SLUG_RE.search(slug) and "thuc-don" not in slug else "menu"
    return {"title": title, "published_at": published.isoformat() if published else None,
            "image_urls": images, "doc_urls": docs, "kind": kind}


def docx_to_text(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = []
            for c in row.cells:  # merged cells repeat; keep one copy
                text = " ".join(c.text.split())
                if not cells or cells[-1] != text:
                    cells.append(text)
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def xlsx_to_text(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    lines = []
    for ws in wb.worksheets[:3]:
        lines.append(f"# sheet {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [" ".join(str(v).split()) for v in row if v is not None and str(v).strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def pdf_to_images(data: bytes) -> list[bytes]:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    pages = []
    for i in range(min(len(pdf), MAX_PDF_PAGES)):
        buf = io.BytesIO()
        pdf[i].render(scale=2).to_pil().save(buf, format="PNG")
        pages.append(buf.getvalue())
    return pages


def load_state() -> dict[tuple[str, str], dict]:
    if not STATE_FILE.exists():
        return {}
    rows = [json.loads(line) for line in STATE_FILE.read_text(encoding="utf-8").splitlines() if line.strip()]
    return {(r["school_code"], r["post_id"]): r for r in rows}


def save_state(state: dict):
    STATE_FILE.parent.mkdir(parents=True, exist_ok=True)
    with STATE_FILE.open("w", encoding="utf-8") as fh:
        for r in sorted(state.values(), key=lambda r: (r["school_code"], r["post_id"])):
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")


async def download_files(f: Fetcher, post: dict) -> list[str]:
    """Save images (and rendered PDF pages) under data/images/<school>/; return local paths."""
    folder = config.IMAGE_DIR / post["school_code"]
    folder.mkdir(parents=True, exist_ok=True)
    paths = []
    for i, url in enumerate(post["image_urls"]):
        data = await f.get_bytes(url)
        if data:
            ext = Path(url.split("?")[0]).suffix.lower() or ".jpg"
            p = folder / f"{post['post_id']}_{i}{ext}"
            p.write_bytes(data)
            paths.append(str(p.relative_to(config.ROOT)))
    for i, url in enumerate(post["doc_urls"]):
        ext = Path(url).suffix.lower()
        if ext in (".doc", ".xls"):
            post["error"] = f"legacy {ext} not supported yet"
            continue
        data = await f.get_bytes(url)
        if not data:
            continue
        try:
            if ext == ".pdf":
                for j, png in enumerate(pdf_to_images(data)):
                    p = folder / f"{post['post_id']}_doc{i}_p{j}.png"
                    p.write_bytes(png)
                    paths.append(str(p.relative_to(config.ROOT)))
            else:
                text = docx_to_text(data) if ext == ".docx" else xlsx_to_text(data)
                p = folder / f"{post['post_id']}_doc{i}.txt"
                p.write_text(text, encoding="utf-8")
                paths.append(str(p.relative_to(config.ROOT)))
        except Exception as e:  # corrupt or encrypted file
            post["error"] = f"{ext} read failed: {e}"
    return paths


async def crawl_school(f: Fetcher, code: str, since: date, state: dict, download: bool) -> list[dict]:
    xml = await f.get_text(site_url(code, "/sitemap.xml"))
    if xml is None:
        print(f"  {code}: sitemap failed")
        return []
    entries, _ = parse_sitemap(xml)
    new = {}
    for url, lastmod in entries:
        if is_menu_post(url) and lastmod and lastmod >= since:
            pid = post_id(url)
            if (code, pid) not in state:
                new.setdefault(pid, (url, lastmod))

    posts = []
    for pid, (url, lastmod) in new.items():
        html = await f.get_text(url)
        if html is None:
            continue
        post = {"school_code": code, "post_id": pid, "url": url, **parse_post(html, url),
                "local_files": [], "status": "pending", "error": None,
                "crawled_at": datetime.now().isoformat(timespec="seconds")}
        post["published_at"] = post["published_at"] or lastmod.isoformat()
        if not post["image_urls"] and not post["doc_urls"]:
            post["status"], post["error"] = "failed", "no images or documents in post"
        elif download:
            post["local_files"] = await download_files(f, post)
            if not post["local_files"]:
                post["status"], post["error"] = "failed", post["error"] or "downloads failed"
        posts.append(post)
    return posts


def pick_schools(status: list[str]) -> list[str]:
    with open(config.DATA_DIR / "coverage.csv", encoding="utf-8") as fh:
        return [r["code"] for r in csv.DictReader(fh) if r["status"] in status]


async def run(codes: list[str], since: date, download: bool):
    state = load_state()
    async with Fetcher(concurrency=2) as f:
        results = await asyncio.gather(*(crawl_school(f, c, since, state, download) for c in codes))
    new_posts = [p for posts in results for p in posts]
    for p in new_posts:
        state[(p["school_code"], p["post_id"])] = p
    save_state(state)
    return new_posts


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--schools", nargs="*", help="school codes (default: from coverage.csv)")
    ap.add_argument("--status", nargs="*", default=["regular", "active"])
    ap.add_argument("--since-days", type=int, default=14)
    ap.add_argument("--no-download", action="store_true")
    args = ap.parse_args()

    codes = args.schools or pick_schools(args.status)
    since = date.today() - timedelta(days=args.since_days)
    print(f"crawling {len(codes)} schools, menu posts since {since}")
    posts = asyncio.run(run(codes, since, not args.no_download))

    by_status = {}
    for p in posts:
        by_status[p["status"]] = by_status.get(p["status"], 0) + 1
    n_doc = sum(1 for p in posts if p["doc_urls"])
    print(f"new posts: {len(posts)} {by_status}; with documents: {n_doc}; tray: {sum(p['kind'] == 'tray' for p in posts)}")
    print(f"state: {STATE_FILE}")


if __name__ == "__main__":
    main()
